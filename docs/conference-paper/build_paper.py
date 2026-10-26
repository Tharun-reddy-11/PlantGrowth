from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).parent
ASSETS = ROOT / "assets"
OUT = ROOT / "Explainable_Temporal_Plant_Growth_Forecasting_Conference_Paper.docx"

TITLE = "An Explainable Adaptive Temporal Learning Framework for Early Plant Growth Forecasting Using Integrated Soil Environmental and Historical Growth Data"
AFFILIATION = "Department of Computer Science and Engineering Artificial Intelligence and Machine Learning\nMohan Babu University, Tirupati, Andhra Pradesh, India"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tcPr.append(shd)
def borders(table):
    tblPr = table._tbl.tblPr; b = OxmlElement("w:tblBorders")
    for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
        x=OxmlElement(f"w:{e}"); x.set(qn("w:val"),"single"); x.set(qn("w:sz"),"4"); x.set(qn("w:color"),"D9D9D9"); b.append(x)
    tblPr.append(b)
def set_cell_text(cell, text, bold=False, color=None, size=8):
    cell.text=""; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0); r=p.add_run(text); r.bold=bold; r.font.size=Pt(size)
    if color:r.font.color.rgb=RGBColor(*color)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
def caption(doc, label):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(5)
    r=p.add_run(label);r.italic=True;r.font.size=Pt(8)
def heading(doc, text):
    p=doc.add_paragraph();p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(2)
    r=p.add_run(text.upper());r.bold=True;r.font.size=Pt(9)
def body(doc, text):
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.0
    for i,part in enumerate(text.split("**")):
        r=p.add_run(part);r.font.size=Pt(8.2);r.bold=(i%2==1)
    return p
def two_cols(section):
    sectPr=section._sectPr; cols=OxmlElement('w:cols');cols.set(qn('w:num'),'2');cols.set(qn('w:space'),'360');sectPr.append(cols)
def make_flow():
    image=Image.new('RGB',(1800,500),'white');draw=ImageDraw.Draw(image);font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',28);small=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',23)
    labels=["Soil\nrecords","Daily\nenvironment","Growth\nhistory","Feature\npipeline","Temporal and\ntabular models"]
    colors=['#DCEFE4','#DCEFE4','#DCEFE4','#CFE5D8','#BEE3D0']
    xs=[45,395,745,1095,1445]
    for label,x,c in zip(labels,xs,colors):
        draw.rounded_rectangle((x,135,x+270,345),radius=22,fill=c,outline='#1E6042',width=5)
        lines=label.split('\n'); y=190
        for line in lines:
            box=draw.textbbox((0,0),line,font=font);draw.text((x+135-(box[2]-box[0])/2,y),line,font=font,fill='#12372A');y+=42
    for x in [320,670,1020,1370]:
        draw.line((x,240,x+60,240),fill='#1E6042',width=6);draw.polygon([(x+60,240),(x+38,228),(x+38,252)],fill='#1E6042')
    draw.text((1290,405),'Forecast + SHAP explanation + chronological evaluation',font=small,fill='#244B36')
    image.save(ASSETS/'framework.png')
def author_block(doc):
    table=doc.add_table(rows=2,cols=3);table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
    entries=[
      ("Ms. B. Rajani", "Project Guide | Assistant Professor\nDepartment of Artificial Intelligence and Machine Learning\nMohan Babu University, Tirupati, Andhra Pradesh, India", ""),
      ("Golla Kalpana Devi", AFFILIATION, "gollakalpanadevi@gmail.com"),
      ("Tammineni Jagadeesh", AFFILIATION, "Jagadeesht805@gmail.com"),
      ("Penumuru Partha", AFFILIATION, ""),
      ("Pathipati Hitesh Naidu", AFFILIATION, "hiteshpathipati45@gmail.com"),
      ("Devarala Vamsi Krishna", AFFILIATION, ""),]
    for cell,(name,aff,email) in zip([c for row in table.rows for c in row.cells],entries):
        cell.width=Inches(2.25);cell.text="";p=cell.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(0)
        r=p.add_run(name+'\n');r.bold=True;r.font.size=Pt(8)
        r=p.add_run(aff+("\n"+email if email else ""));r.font.size=Pt(6.7)
    for row in table.rows:
        for c in row.cells:
            tcPr=c._tc.get_or_add_tcPr();tcBorders=OxmlElement('w:tcBorders');tcPr.append(tcBorders)
def main():
    make_flow();doc=Document();sec=doc.sections[0];sec.top_margin=Inches(.55);sec.bottom_margin=Inches(.55);sec.left_margin=Inches(.55);sec.right_margin=Inches(.55)
    styles=doc.styles;styles['Normal'].font.name='Times New Roman';styles['Normal'].font.size=Pt(8.2)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(5)
    r=p.add_run(TITLE);r.bold=True;r.font.size=Pt(15);r.font.name='Times New Roman'
    author_block(doc)
    p=doc.add_paragraph();p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(4)
    r=p.add_run('Abstract - ');r.bold=True;r.font.size=Pt(8.5)
    r=p.add_run('Plant growth reflects interacting soil conditions, environmental exposure, and prior development. This paper presents an explainable adaptive temporal learning framework for forecasting future plant height from integrated soil, daily environmental, and historical growth observations. The framework constructs leakage-safe lagged and rolling features, benchmarks persistence, linear, Random Forest, XGBoost, GRU, and LSTM models, and uses a validation-trained fusion layer to combine tabular and temporal signals. Feature selection is fitted on the training period only, while SHAP provides model-based local and global explanations. Evaluation is designed around chronological train, validation, and held-out test partitions using MAE, RMSE, and R squared. The implementation is delivered as a React, Spring Boot, MySQL, and FastAPI architecture in which the backend persists prediction inputs, model metadata, and explanations. The current supplied dataset inventory contains environmental and aggregate greenhouse measurements but no repeated plant-height target; consequently, this paper reports a reproducible methodology and data-quality observation rather than fabricated accuracy values.');r.font.size=Pt(8.2)
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(5);r=p.add_run('Index Terms - ');r.bold=True;r.font.size=Pt(8.2);r=p.add_run('plant growth forecasting, temporal learning, explainable artificial intelligence, SHAP, feature selection, precision agriculture.');r.font.size=Pt(8.2)
    section=doc.add_section(WD_SECTION.CONTINUOUS);two_cols(section)
    heading(doc,'I. Introduction')
    body(doc,'Plant growth monitoring is central to data-driven cultivation, yet manual measurement alone may not expose the temporal effects of weather, soil, and prior development. A useful forecasting system must separate measured observations from future estimates and avoid treating later measurements as input evidence. This work targets future plant height in centimetres at a declared forecast horizon.')
    body(doc,'The proposed framework is designed for research use: it records provenance, evaluates models chronologically, and returns an explanation associated with the exact fitted artifact and input snapshot. It does not diagnose plant disease, guarantee yield, or infer causal biological effects from model attributions.')
    heading(doc,'II. Related Work and Research Gap')
    body(doc,'Machine-learning approaches have been applied to greenhouse growth modelling and environmental response prediction. Cha et al. used LSTM modelling for tomato plant height with environmental determinants. Riesselman et al. demonstrated multimodal trajectory prediction from industrial greenhouse observations. Explainable agricultural modelling has also used SHAP to examine feature contributions. These directions motivate temporal and multimodal forecasting, but they do not by themselves provide a unified, auditable application that combines soil records, environmental observations, historical height, feature selection, fusion, and chronological evaluation.')
    body(doc,'The FBG wearable sensor study by Condò et al. is a distinct base-paper context. It measures stem elongation through a self-adhesive optical sensor. In contrast, the present system forecasts a later numerical growth value from historical and contextual records. The two approaches may be complementary, but this software neither reproduces nor replaces the sensor measurement.')
    heading(doc,'III. Proposed Framework')
    doc.add_picture(str(ASSETS/'framework.png'),width=Inches(3.45));caption(doc,'Fig. 1. Proposed integrated forecasting and explanation workflow.')
    body(doc,'At each forecast origin, the pipeline joins records by plant, location, and timestamp; creates only features available at that origin; and stores a versioned feature snapshot. Tabular models consume selected current, lagged, rolling, and static features. GRU and LSTM models consume per-plant ordered sequences. A stacking fusion model is trained only from temporally valid base-model predictions. The deployed service rejects a request when a required artifact or compatible horizon is absent.')
    heading(doc,'IV. Data and Feature Design')
    table=doc.add_table(rows=1,cols=3);table.alignment=WD_TABLE_ALIGNMENT.CENTER;borders(table)
    for c,t in zip(table.rows[0].cells,['Modality','Examples','Integrity rule']):shade(c,'1E6042');set_cell_text(c,t,True,(255,255,255),7.5)
    for row in [('Soil','pH, moisture, N, P, K, texture','Keep unknown values null; record source and depth.'),('Environment','temperature, humidity, precipitation, solar radiation','Keep original units; do not relabel radiation as lux.'),('Growth','timestamped plant height, growth stage','No fictional daily values; preserve plant identity and time.')]:
        cells=table.add_row().cells
        for c,t in zip(cells,row):set_cell_text(c,t,size=7.2)
    caption(doc,'Table I. Input modalities and integrity controls.')
    body(doc,'Temporal features include previous height, observed growth rate, rolling temperature and humidity summaries, cumulative precipitation, intervals between observations, and soil attributes where their temporal meaning permits reuse. Missing-value transforms are fitted from training data only. A valid future-height label is the nearest real measurement within a documented tolerance of the requested horizon; unmatched origins are excluded.')
    heading(doc,'V. Models and Explainability')
    body(doc,'The benchmark includes a persistence baseline and linear regression, followed by Random Forest and XGBoost for nonlinear tabular learning. GRU and LSTM represent ordered per-plant sequences. The primary fusion strategy is validation-trained stacking: base-model predictions generated without future leakage form the meta-regressor inputs. Adaptive feature selection uses training-period model importance or permutation importance, with selected feature names stored with the artifact.')
    body(doc,'SHAP explanations are calculated for the fitted model and the exact transformed input. Tree explainers are used for tree models; temporal models require a validated neural explainer and representative training background. SHAP values describe model contributions, not causal effects.')
    heading(doc,'VI. Evaluation Protocol and Current Observation')
    table=doc.add_table(rows=1,cols=3);table.alignment=WD_TABLE_ALIGNMENT.CENTER;borders(table)
    for c,t in zip(table.rows[0].cells,['Stage','Purpose','Leakage control']):shade(c,'1E6042');set_cell_text(c,t,True,(255,255,255),7.5)
    for row in [('Train: earliest period','Fit imputers, scalers, feature selector, models','No validation/test labels used.'),('Validation: following period','Select hyperparameters and fusion configuration','No final test tuning.'),('Test: newest period','One-time held-out assessment','Report MAE, RMSE, R squared, dates, sample count.')]:
        cells=table.add_row().cells
        for c,t in zip(cells,row):set_cell_text(c,t,size=7.2)
    caption(doc,'Table II. Chronological evaluation design.')
    body(doc,'Inspection of the currently supplied aggregate greenhouse file found that its plant-height field is blank and explicitly marked as unavailable. It cannot support a plant-height experiment or a valid seven-day forecast evaluation. This is a data-quality finding, not a negative model result. Reported metrics and comparative claims must be added only after repeated, timestamped plant-height observations are collected or licensed from a compatible source.')
    heading(doc,'VII. System Implementation')
    body(doc,'The application uses a React client, Spring Boot REST API, MySQL persistence layer, and separate FastAPI machine-learning service. Spring Boot authenticates users using JWT, enforces plant ownership, retrieves authorized records, and persists the request snapshot and response metadata. The Python service owns preprocessing, artifact loading, prediction, evaluation, and explanation. This separation prevents the browser from directly accessing the model service and supports reproducibility.')
    heading(doc,'VIII. Conclusion and Future Work')
    body(doc,'This paper specifies an explainable temporal framework for early plant-growth forecasting that integrates soil, environmental, and historical measurements while preserving temporal validity. The immediate next step is to acquire repeated height observations at an adequate sampling frequency, train every benchmark for supported horizons, and publish chronological held-out results. Future work may compare location-specific and sensor-assisted data streams, quantify uncertainty, and evaluate generalization across species and growing environments.')
    heading(doc,'References')
    refs=[
      '[1] L. Cha et al., “Artificial intelligence based prediction of tomato plant height using environmental determinants,” Computers and Electronics in Agriculture, 2022.',
      '[2] M. Riesselman et al., “Self-supervised learning of plant growth trajectories from multimodal greenhouse data,” 2024.',
      '[3] I. Condò et al., “An innovative self-adhesive wearable sensor based on fiber Bragg grating technology for plant growth monitoring,” IEEE Sensors Journal, vol. 25, no. 2, pp. 2705-2713, 2025.',
      '[4] S. M. Lundberg and S.-I. Lee, “A unified approach to interpreting model predictions,” Advances in Neural Information Processing Systems, 2017.',
      '[5] ISRIC, “SoilGrids,” https://soilgrids.org/.',
      '[6] NASA POWER, “Daily API,” https://power.larc.nasa.gov/docs/services/api/temporal/daily/.',
    ]
    for ref in refs: body(doc,ref)
    # Footer page number field
    for s in doc.sections:
        p=s.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run('Page ');fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld);r.font.size=Pt(8)
    doc.core_properties.title=TITLE;doc.core_properties.author='Golla Kalpana Devi; Tammineni Jagadeesh; Penumuru Partha; Pathipati Hitesh Naidu; Devarala Vamsi Krishna'
    doc.save(OUT)
if __name__=='__main__': main()
